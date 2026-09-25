#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera a planilha de análise do estoque no Mercado Livre Full.

Cruza o relatório "stock_general_full" do ML com o custo e o preço de venda
líquido do catálogo (JSONs dos itens exportados do app Estoque Vivo).

Uso:
    python3 scripts/analise_full_ml.py RELATORIO_ML.xlsx DIR_ITENS SAIDA.xlsx DD/MM/AAAA

Os valores de cada fórmula são calculados aqui, na mesma passada que escreve a
fórmula, e gravados como cache no XML: o LibreOffice headless não roda neste
ambiente, então não há recálculo externo. O arquivo sai com fullCalcOnLoad=1,
e o Excel/Sheets recalcula tudo ao abrir.
"""
import glob
import json
import os
import re
import shutil
import sys
import zipfile
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# SKU do anúncio no ML -> id do item no app
MAPA = {
    "1942full": "ret-1942", "5702full": "ret-5702", "9203full": "ved-9203", "2400full": "ret-2400",
    "1035jfull": "bronzina-1035j", "2317": "ret-2317", "3044full": "ret-3044", "2283full": "ret-2283",
    "5266": "ret-5266", "7092": "anel-7092", "5245full": "ret-5245", "5772strada": "ret-5772",
    "2371full": "ret-2371", "5159full": "ret-5159", "2539": "ret-2539", "8126full": "anel-8126-std",
    "5338full": "ret-5338", "2178full": "ret-2178", "5502": "ret-5502", "2525full": "ret-2525",
    "2544full": "ret-2544", "1135": "ret-1135", "2075": "ret-2075", "5601full": "ret-5601",
    "5699full": "ret-5699", "2370full": "ret-2370", "7224STDfull": "jogo-7224-std", "5801full": "ret-5801",
    "2373full": "ret-2373", "6631050full": "anel-6631-050", "2374full": "ret-2374",
}

# colunas da aba "Resumo" do relatório do ML, conferidas pelo cabeçalho antes de ler
CABECALHO_ML = {"B2": "SKU", "C2": "Produto", "D3": "Entrada pendente", "E3": "Em transferência",
                "F3": "Aptas para venda", "K3": "Entrada pendente", "N3": "Vendas\núlt. 30 dias (R$)",
                "O3": "Unidades vendidas\núlt. 30 días", "P3": "Urgência de envio"}

ARIAL, BRANCO, MARCA, AMARELO, AZUL, CINZA, SUAVE = (
    "Arial", "FFFFFFFF", "FF0C5B6A", "FFFFF3C4", "FF0000FF", "FFF2F5F6", "FF5D7078")
MOEDA = 'R$ #,##0.00;-R$ #,##0.00;"—"'
BORDA = Border(bottom=Side(style="thin", color="FFD6E0E4"))


def fonte(bold=False, cor=None, size=10, italic=False):
    return Font(name=ARIAL, bold=bold, size=size, italic=italic, **({"color": cor} if cor else {}))


def ler_relatorio(caminho):
    ws = load_workbook(caminho, data_only=True)["Resumo"]
    for ref, esperado in CABECALHO_ML.items():
        if (ws[ref].value or "").strip() != esperado.strip():
            sys.exit("Formato do relatório mudou: %s = %r, esperado %r" % (ref, ws[ref].value, esperado))
    linhas, r = [], 5
    while ws["B%d" % r].value:
        g = lambda c: ws["%s%d" % (c, r)].value
        linhas.append(dict(sku_ml=g("B"), produto=g("C"), pendente=g("D"), transf=g("E"), aptas=g("F"),
                           nao=g("G") + g("H") + g("I"), pend_espaco=g("K"), vendas_rs=g("N"),
                           vendas_un=g("O"), urgencia=g("P"), a_enviar=g("Q")))
        r += 1
    total_ml = {"J": ws["J%d" % r].value, "K": ws["K%d" % r].value}
    soma_k = sum(l["pend_espaco"] for l in linhas)
    if soma_k != total_ml["K"]:
        sys.exit("Soma da coluna K (%d) não bate com o total do ML (%d)" % (soma_k, total_ml["K"]))
    sem_mapa = [l["sku_ml"] for l in linhas if l["sku_ml"] not in MAPA]
    if sem_mapa:
        sys.exit("SKUs do ML sem correspondência no catálogo: %s" % sem_mapa)
    return linhas


def gerar(relatorio, dir_itens, saida, data_txt):
    linhas = ler_relatorio(relatorio)
    cat = {os.path.basename(f)[:-5]: json.load(open(f, encoding="utf-8"))
           for f in glob.glob(os.path.join(dir_itens, "*.json"))}

    reg = []
    for l in linhas:
        it = cat[MAPA[l["sku_ml"]]]
        x = dict(l, codigo=it.get("sku", ""), custo=it.get("custo") or 0, preco=it.get("preco") or 0,
                 galpao=it.get("qtd") or 0, devol=l["pend_espaco"] - l["pendente"])
        x["total"] = x["aptas"] + x["transf"] + x["nao"]
        reg.append(x)
    reg.sort(key=lambda x: (-x["total"] * x["custo"], x["sku_ml"]))

    galpao = dict(un=sum(it.get("qtd") or 0 for it in cat.values()),
                  custo=round(sum((it.get("qtd") or 0) * (it.get("custo") or 0) for it in cat.values()), 2),
                  rec=round(sum((it.get("qtd") or 0) * (it.get("preco") or 0) for it in cat.values()), 2))
    sem_preco_galpao = sorted(it["sku"] for it in cat.values() if not it.get("preco"))

    cache = {"Resumo": {}, "Full por SKU": {}}
    wb = Workbook()

    # ------------------------------ Full por SKU ------------------------------
    det = wb.active
    det.title = "Full por SKU"
    C = dict(sku=1, cod=2, prod=3, apt=4, trf=5, nao=6, tot=7, cam=8, dev=9, cus=10, pre=11, cpar=12,
             rec=13, mar=14, vun=15, vrs=16, tkt=17, cob=18, urg=19, env=20, gal=21)
    L = {k: get_column_letter(v) for k, v in C.items()}
    titulos = [("SKU no ML", 13), ("Código", 11), ("Produto", 44), ("Aptas para venda", 10),
               ("Em transferência", 11), ("Não aptas", 8), ("Total no Full", 9), ("A caminho do Full", 10),
               ("Devoluções voltando", 11), ("Custo unit. (R$)", 11), ("Preço venda líquido (R$)", 11),
               ("Custo parado no Full (R$)", 14), ("Receita projetada (R$)", 14), ("Margem projetada (R$)", 14),
               ("Vendidas 30 dias", 9), ("Vendas 30 dias no ML (R$)", 14), ("Ticket médio no ML (R$)", 11),
               ("Cobertura (dias)", 10), ("Urgência de envio (ML)", 14), ("ML sugere enviar", 9),
               ("No seu galpão (app)", 10)]
    for c, (t, w) in enumerate(titulos, 1):
        cel = det.cell(row=1, column=c, value=t)
        cel.font, cel.fill = fonte(True, BRANCO), PatternFill("solid", fgColor=MARCA)
        cel.alignment = Alignment(vertical="center", wrap_text=True)
        det.column_dimensions[get_column_letter(c)].width = w
    det.row_dimensions[1].height = 44
    det.freeze_panes = "D2"

    L0 = 2
    ULT = L0 + len(reg) - 1
    TOT = ULT + 1
    inteiros = {C[k] for k in ("apt", "trf", "nao", "tot", "cam", "dev", "vun", "env", "gal")}
    dinheiro = {C[k] for k in ("cus", "pre", "cpar", "rec", "mar", "vrs", "tkt")}
    cd = cache["Full por SKU"]

    def cob(apt, trf, vun):
        return round((apt + trf) / (vun / 30.0), 1) if vun else ""

    for i, x in enumerate(reg):
        r = L0 + i
        for k, v in (("sku", x["sku_ml"]), ("cod", x["codigo"]), ("prod", x["produto"]), ("apt", x["aptas"]),
                     ("trf", x["transf"]), ("nao", x["nao"]), ("cam", x["pendente"]), ("dev", x["devol"]),
                     ("cus", x["custo"]), ("pre", x["preco"]), ("vun", x["vendas_un"]), ("vrs", x["vendas_rs"]),
                     ("urg", x["urgencia"]), ("env", x["a_enviar"]), ("gal", x["galpao"])):
            det.cell(row=r, column=C[k], value=v)
        f = lambda k, frm, val: (det.cell(row=r, column=C[k], value=frm), cd.__setitem__("%s%d" % (L[k], r), val))
        f("tot", "=%s%d+%s%d+%s%d" % (L["apt"], r, L["trf"], r, L["nao"], r), x["total"])
        f("cpar", "=%s%d*%s%d" % (L["tot"], r, L["cus"], r), round(x["total"] * x["custo"], 2))
        f("rec", "=%s%d*%s%d" % (L["tot"], r, L["pre"], r), round(x["total"] * x["preco"], 2))
        f("mar", "=%s%d-%s%d" % (L["rec"], r, L["cpar"], r), round(x["total"] * (x["preco"] - x["custo"]), 2))
        f("tkt", "=IF({v}{r}>0,{s}{r}/{v}{r},0)".format(v=L["vun"], s=L["vrs"], r=r),
          x["vendas_rs"] / x["vendas_un"] if x["vendas_un"] else 0)
        f("cob", '=IF({v}{r}>0,ROUND(({a}{r}+{t}{r})/({v}{r}/30),1),"")'.format(v=L["vun"], a=L["apt"], t=L["trf"], r=r),
          cob(x["aptas"], x["transf"], x["vendas_un"]))
        for col in range(1, len(titulos) + 1):
            cel = det.cell(row=r, column=col)
            cel.font, cel.border = fonte(), BORDA
            if col in inteiros: cel.number_format = "#,##0"
            if col in dinheiro: cel.number_format = MOEDA
            if col == C["cob"]: cel.number_format = "0.0"
        for k in ("cus", "pre"):
            det.cell(row=r, column=C[k]).fill = PatternFill("solid", fgColor=AMARELO)
            det.cell(row=r, column=C[k]).font = fonte(cor=AZUL)

    S = lambda k: sum(x[k] for x in reg)
    det.cell(row=TOT, column=1, value="TOTAL")
    somas = {"apt": S("aptas"), "trf": S("transf"), "nao": S("nao"), "tot": S("total"), "cam": S("pendente"),
             "dev": S("devol"), "vun": S("vendas_un"), "vrs": round(S("vendas_rs"), 2), "env": S("a_enviar"),
             "gal": S("galpao")}
    for k in ("cpar", "rec", "mar"):
        somas[k] = round(sum(cd["%s%d" % (L[k], r)] for r in range(L0, ULT + 1)), 2)
    for k, v in somas.items():
        det.cell(row=TOT, column=C[k], value="=SUM({c}{a}:{c}{b})".format(c=L[k], a=L0, b=ULT))
        cd["%s%d" % (L[k], TOT)] = v
    det.cell(row=TOT, column=C["tkt"], value="=IF({v}{r}>0,{s}{r}/{v}{r},0)".format(v=L["vun"], s=L["vrs"], r=TOT))
    cd["%s%d" % (L["tkt"], TOT)] = somas["vrs"] / somas["vun"]
    det.cell(row=TOT, column=C["cob"],
             value='=IF({v}{r}>0,ROUND(({a}{r}+{t}{r})/({v}{r}/30),1),"")'.format(v=L["vun"], a=L["apt"], t=L["trf"], r=TOT))
    cd["%s%d" % (L["cob"], TOT)] = cob(somas["apt"], somas["trf"], somas["vun"])
    for col in range(1, len(titulos) + 1):
        cel = det.cell(row=TOT, column=col)
        cel.font, cel.fill = fonte(True), PatternFill("solid", fgColor=CINZA)
        cel.border = Border(top=Side(style="medium", color=MARCA))
        if col in inteiros: cel.number_format = "#,##0"
        if col in dinheiro: cel.number_format = MOEDA
        if col == C["cob"]: cel.number_format = "0.0"

    det.auto_filter.ref = "A1:%s%d" % (L["gal"], ULT)
    vermelho = dict(fill=PatternFill("solid", fgColor="FFFAE6E4"), font=Font(name=ARIAL, size=10, bold=True, color="FFAF2B23"))
    laranja = dict(fill=PatternFill("solid", fgColor="FFFBEEDA"), font=Font(name=ARIAL, size=10, bold=True, color="FF9D5C05"))
    fx_cob = "%s2:%s%d" % (L["cob"], L["cob"], ULT)
    fx_urg = "%s2:%s%d" % (L["urg"], L["urg"], ULT)
    det.conditional_formatting.add(fx_cob, CellIsRule(operator="lessThan", formula=["7"], **vermelho))
    det.conditional_formatting.add(fx_cob, CellIsRule(operator="greaterThan", formula=["90"], **laranja))
    det.conditional_formatting.add(fx_urg, CellIsRule(operator="equal", formula=['"Urgente"'], **vermelho))
    for st in ("Esta semana", "Próxima semana"):
        det.conditional_formatting.add(fx_urg, CellIsRule(operator="equal", formula=['"%s"' % st], **laranja))
    det.sheet_view.showGridLines = False

    # ------------------------------ Resumo ------------------------------
    res = wb.create_sheet("Resumo", 0)
    cr = cache["Resumo"]
    for col, w in zip("ABCDEF", (40, 16, 18, 20, 18, 60)):
        res.column_dimensions[col].width = w
    res["A1"] = "Mercado Livre Full — posição de %s" % data_txt
    res["A1"].font = fonte(True, MARCA, 15)
    res["A2"] = "Fonte: relatório de estoque Full do Mercado Livre + custo e preço de venda líquido do seu catálogo (app Estoque Vivo)."
    res["A2"].font = fonte(cor=SUAVE, italic=True)

    D = "'Full por SKU'!"
    rng = lambda k: "%s%s%d:%s%d" % (D, L[k], L0, L[k], ULT)

    def estilo(r, negrito=False, cinza=False, fmt_b="#,##0"):
        for col in range(1, 7):
            cel = res.cell(row=r, column=col)
            cel.font = fonte(negrito, cor=(SUAVE if col == 6 else None), size=(9 if col == 6 else 10))
            cel.border = BORDA
            if cinza: cel.fill = PatternFill("solid", fgColor=CINZA)
            if col == 6: cel.alignment = Alignment(wrap_text=True, vertical="top")
        res.cell(row=r, column=2).number_format = fmt_b
        for col in (3, 4, 5):
            res.cell(row=r, column=col).number_format = MOEDA

    def cab(r, titulos_):
        for i, t in enumerate(titulos_):
            cel = res.cell(row=r, column=1 + i, value=t)
            cel.font, cel.fill = fonte(True, BRANCO), PatternFill("solid", fgColor=MARCA)
            cel.alignment = Alignment(vertical="center", wrap_text=True)
        res.row_dimensions[r].height = 30

    def secao(r, texto):
        res.cell(row=r, column=1, value=texto).font = fonte(True, MARCA, 12)

    def put(ref, formula, valor):
        res[ref] = formula
        cr[ref] = valor

    def linha_status(r, rotulo, k, campo, obs, negrito=False):
        un = S(campo)
        custo = round(sum(x[campo] * x["custo"] for x in reg), 2)
        rec = round(sum(x[campo] * x["preco"] for x in reg), 2)
        res.cell(row=r, column=1, value=rotulo)
        put("B%d" % r, "=SUM(%s)" % rng(k), un)
        put("C%d" % r, "=SUMPRODUCT(%s,%s)" % (rng(k), rng("cus")), custo)
        put("D%d" % r, "=SUMPRODUCT(%s,%s)" % (rng(k), rng("pre")), rec)
        put("E%d" % r, "=D%d-C%d" % (r, r), round(rec - custo, 2))
        res.cell(row=r, column=6, value=obs)
        estilo(r, negrito)

    cab(4, ["Situação no Mercado Livre", "Unidades", "Custo parado (R$)", "Receita projetada (R$)",
            "Margem projetada (R$)", "O que significa"])
    linha_status(5, "Aptas para venda", "apt", "aptas", "Já estão no Full e à venda.")
    linha_status(6, "Em transferência entre galpões do ML", "trf", "transf",
                 "Dentro do Full, mudando de um galpão do ML para outro.")
    linha_status(7, "Temporariamente não aptas", "nao", "nao", "Extraviadas, em revisão ou de vendas canceladas.")
    res["A8"] = "TOTAL NO FULL"
    for col in "BCDE":
        put("%s8" % col, "=SUM(%s5:%s7)" % (col, col), round(sum(cr["%s%d" % (col, r)] for r in (5, 6, 7)), 2))
    res["F8"] = "Dinheiro parado dentro dos galpões do Mercado Livre."
    estilo(8, True, True)
    linha_status(9, "A caminho do Full (entrada pendente)", "cam", "pendente",
                 "Envios que o ML ainda não recebeu. Saíram do galpão antes da contagem de 18/09, então o app já não conta essas peças.")
    linha_status(10, "Devoluções voltando ao Full", "dev", "devol",
                 "Vendas devolvidas a caminho do Full. Podem voltar avariadas: o valor é o máximo.")

    secao(12, "Visão da empresa")
    cab(13, ["Onde está", "Unidades", "Custo parado (R$)", "Receita projetada (R$)", "Margem projetada (R$)", "Observação"])
    res["A14"] = "Seu galpão (app Estoque Vivo)"
    res["B14"], res["C14"], res["D14"] = galpao["un"], galpao["custo"], galpao["rec"]
    put("E14", "=D14-C14", round(galpao["rec"] - galpao["custo"], 2))
    res["F14"] = "Números do app em %s.%s" % (
        data_txt, (" Sem preço de venda: %s." % ", ".join(sem_preco_galpao)) if sem_preco_galpao else "")
    estilo(14)
    for col in "BCD":
        res["%s14" % col].font = fonte(cor=AZUL)
        res["%s14" % col].comment = Comment("Número copiado do app Estoque Vivo (não é fórmula).", "Estoque Vivo")
    for r, rot, origem, obs in ((15, "Dentro do Full", 8, "Total no Full, da tabela acima."),
                                (16, "A caminho do Full", 9, ""),
                                (17, "Devoluções voltando ao Full", 10, "")):
        res["A%d" % r] = rot
        for col in "BCDE":
            put("%s%d" % (col, r), "=%s%d" % (col, origem), cr["%s%d" % (col, origem)])
        res["F%d" % r] = obs
        estilo(r)
    res["A18"] = "TOTAL DA EMPRESA"
    for col in "BCDE":
        put("%s18" % col, "=SUM(%s14:%s17)" % (col, col),
            round((galpao["un"] if col == "B" else galpao["custo"] if col == "C" else galpao["rec"] if col == "D"
                   else galpao["rec"] - galpao["custo"]) + sum(cr["%s%d" % (col, r)] for r in (15, 16, 17)), 2))
    res["F18"] = "Tudo o que é seu: galpão, Full, a caminho e devoluções."
    estilo(18, True, True)

    secao(20, "Referência do Mercado Livre — últimos 30 dias")
    vrs, vun = somas["vrs"], somas["vun"]
    liquido_30d = sum(x["vendas_un"] * x["preco"] for x in reg)
    ticket_bruto_full = round(sum(x["total"] * (x["vendas_rs"] / x["vendas_un"] if x["vendas_un"] else 0) for x in reg), 2)
    refs = [
        (21, "Vendas no ML (R$, bruto)", "=SUM(%s)" % rng("vrs"), vrs, MOEDA, "Valor de venda registrado pelo ML, antes das tarifas."),
        (22, "Unidades vendidas", "=SUM(%s)" % rng("vun"), vun, "#,##0", ""),
        (23, "Ticket médio no ML (R$)", "=IF(B22>0,B21/B22,0)", vrs / vun, MOEDA, "Preço médio de venda no ML por unidade."),
        (24, "De cada R$ 100 vendidos no ML, você recebe", "=IF(B21>0,SUMPRODUCT(%s,%s)/B21*100,0)" % (rng("vun"), rng("pre")),
         liquido_30d / vrs * 100, MOEDA, "Estimativa: unidades vendidas × seu preço líquido ÷ vendas brutas. O resto são tarifas e frete do ML."),
        (25, "Estoque do Full pelo ticket médio do ML (R$, bruto)", "=SUMPRODUCT(%s,%s)" % (rng("tot"), rng("tkt")),
         ticket_bruto_full, MOEDA, "Quanto o estoque do Full faturaria no ML pelo preço médio do último mês, antes das tarifas."),
        (26, "Cobertura média do Full (dias)", "=IF(B22>0,ROUND((B5+B6)/(B22/30),1),0)",
         cob(somas["apt"], somas["trf"], vun), "0.0", "Quantos dias o estoque do Full dura no ritmo de venda atual."),
    ]
    for r, rot, frm, v, fmt, obs in refs:
        res.cell(row=r, column=1, value=rot)
        put("B%d" % r, frm, v)
        res.cell(row=r, column=6, value=obs)
        estilo(r, fmt_b=fmt)

    secao(28, "Alertas de reposição (aba Full por SKU)")
    cobs = [cob(x["aptas"], x["transf"], x["vendas_un"]) for x in reg]
    alertas = [
        (29, 'Urgência "Urgente"', '=COUNTIF(%s,"Urgente")' % rng("urg"), sum(x["urgencia"] == "Urgente" for x in reg),
         "Status que o próprio ML calcula para cada anúncio."),
        (30, 'Urgência "Esta semana"', '=COUNTIF(%s,"Esta semana")' % rng("urg"), sum(x["urgencia"] == "Esta semana" for x in reg), ""),
        (31, 'Urgência "Próxima semana"', '=COUNTIF(%s,"Próxima semana")' % rng("urg"), sum(x["urgencia"] == "Próxima semana" for x in reg), ""),
        (32, 'Estoque "Excedente" no Full', '=COUNTIF(%s,"Excedente")' % rng("urg"), sum(x["urgencia"] == "Excedente" for x in reg),
         "O ML considera que há estoque demais. Pode gerar tarifa de armazenagem."),
        (33, "SKUs com menos de 7 dias de cobertura", '=COUNTIF(%s,"<7")' % rng("cob"), sum(1 for v in cobs if v != "" and v < 7), ""),
    ]
    for r, rot, frm, v, obs in alertas:
        res.cell(row=r, column=1, value=rot)
        put("B%d" % r, frm, v)
        res.cell(row=r, column=6, value=obs)
        estilo(r)

    secao(35, "Como ler")
    notas = [
        "Receita projetada = unidades × seu preço de venda líquido (o que você recebe por peça, já sem as tarifas do ML). Custo parado = unidades × custo. Margem = receita − custo.",
        "Custo e preço vêm do seu catálogo (células amarelas da aba Full por SKU). Se mudar um valor ali, o resumo recalcula sozinho.",
        "A caminho do Full: saíram do galpão antes da contagem de 18/09. O 5801 e o 5245 confirmam isso: o galpão nunca teve essas quantidades depois da contagem. Por isso somam no total da empresa sem contar duas vezes.",
        "Devoluções: o ML conta %d unidades em entrada pendente na coluna de espaço ocupado e %d na coluna \"a caminho\". A diferença (%d) são devoluções de clientes voltando ao Full." % (
            sum(x["pend_espaco"] for x in reg), somas["cam"], somas["dev"]),
        "Cobertura = (aptas + em transferência) ÷ venda diária dos últimos 30 dias. Vermelho abaixo de 7 dias, laranja acima de 90.",
    ]
    r = 36
    for t in notas:
        cel = res.cell(row=r, column=1, value="• " + t)
        cel.font, cel.alignment = fonte(), Alignment(wrap_text=True, vertical="top")
        res.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        res.row_dimensions[r].height = 30
        r += 1
    res.sheet_view.showGridLines = False

    wb.save(saida)
    injetar_cache(saida, cache)
    validar(saida, cache)
    return cache


CELL_RE = re.compile(r'<c\b[^>]*\br="(?P<ref>[A-Z]+\d+)"[^>]*?(?:/>|>.*?</c>)', re.DOTALL)


def injetar_cache(caminho, cache):
    with zipfile.ZipFile(caminho) as z:
        wbx = z.read("xl/workbook.xml").decode("utf-8")
        rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    alvo = {rid: t for t, rid in re.findall(r'Target="/?(xl/worksheets/sheet\d+\.xml)" Id="(rId\d+)"', rels)}
    arquivo = {alvo[rid]: nome for nome, rid in re.findall(r'<sheet name="([^"]+)" sheetId="\d+" state="visible" r:id="(rId\d+)"', wbx)}

    def inject(texto, mapa):
        def repl(m):
            full, ref = m.group(0), m.group("ref")
            if ref not in mapa or "<f" not in full:
                return full
            v = mapa[ref]
            if isinstance(v, str):
                novo = full if re.search(r"<c\b[^>]*\bt=", full) else re.sub(r"^<c\b", '<c t="str"', full, count=1)
                vx = escape(v)
            else:
                novo, vx = full, repr(float(v)) if isinstance(v, float) else str(v)
            if re.search(r"<v\s*/>|<v>\s*</v>", novo):
                return re.sub(r"<v\s*/>|<v>\s*</v>", "<v>%s</v>" % vx, novo, count=1)
            return novo if "<v>" in novo else novo.replace("</f>", "</f><v>%s</v>" % vx, 1)
        return CELL_RE.sub(repl, texto)

    tmp = caminho + ".tmp"
    with zipfile.ZipFile(caminho) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in arquivo:
                data = inject(data.decode("utf-8"), cache[arquivo[item.filename]]).encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, caminho)


def validar(caminho, cache):
    wb, wbf = load_workbook(caminho, data_only=True), load_workbook(caminho)
    erros = []
    for nome, mapa in cache.items():
        for ref, esp in mapa.items():
            lido = wb[nome][ref].value
            if not (lido == esp or (isinstance(esp, (int, float)) and isinstance(lido, (int, float)) and abs(lido - esp) < 1e-6)):
                erros.append("%s!%s esperado=%r lido=%r" % (nome, ref, esp, lido))
    for ws in wbf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and c.coordinate not in cache[ws.title]:
                    erros.append("fórmula sem valor calculado: %s!%s" % (ws.title, c.coordinate))
    with zipfile.ZipFile(caminho) as z:
        for n in z.namelist():
            if n.endswith(".xml"):
                ET.fromstring(z.read(n))
    if erros:
        sys.exit("Validação falhou:\n  " + "\n  ".join(erros))
    print("validado: %d fórmulas com valor calculado, 0 divergências, XML bem formado" % sum(len(m) for m in cache.values()))


if __name__ == "__main__":
    if len(sys.argv) != 5:
        sys.exit(__doc__)
    c = gerar(*sys.argv[1:])
    r = c["Resumo"]
    print("Full:        %5d un | custo R$ %10.2f | receita R$ %10.2f" % (r["B8"], r["C8"], r["D8"]))
    print("A caminho:   %5d un | custo R$ %10.2f | receita R$ %10.2f" % (r["B9"], r["C9"], r["D9"]))
    print("Devoluções:  %5d un | custo R$ %10.2f | receita R$ %10.2f" % (r["B10"], r["C10"], r["D10"]))
    print("Empresa:     %5d un | custo R$ %10.2f | receita R$ %10.2f" % (r["B18"], r["C18"], r["D18"]))
    print("De cada R$ 100 no ML você recebe: R$ %.2f" % r["B24"])
