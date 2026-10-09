#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera a planilha do estoque no Mercado Livre Full a partir de um print do relatório.

Quando só há o print da tela do ML (sem o .xlsx do relatório), os números são
transcritos para um CSV com as colunas:
    sku_ml, produto, entrada_pendente, em_transferencia, devolvidas, aptas,
    ocupam_espaco, distribuidas_entrada_pendente
O print não traz vendas nem urgência, então a planilha sai sem essas colunas.
"Não aptas" é o que sobra de "ocupam espaço" depois das outras colunas (o print
esconde as colunas de não aptas do ML).

Uso:
    python3 scripts/full_ml_print.py TRANSCRICAO.csv DIR_ITENS SAIDA.xlsx DD/MM/AAAA
"""
import csv
import glob
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from analise_full_ml import (AMARELO, AZUL, BORDA, BRANCO, CINZA, MAPA, MARCA, MOEDA, SUAVE,  # noqa: E402
                             fonte, injetar_cache, validar)


def ler_transcricao(caminho):
    linhas = []
    with open(caminho, encoding="utf-8") as f:
        for r in csv.DictReader(f):
            x = {k: (v if k in ("sku_ml", "produto") else int(v)) for k, v in r.items()}
            x["nao"] = x["ocupam_espaco"] - x["entrada_pendente"] - x["em_transferencia"] - x["devolvidas"] - x["aptas"]
            if x["nao"] < 0:
                sys.exit("%s: colunas somam mais que 'ocupam espaço'" % x["sku_ml"])
            if x["distribuidas_entrada_pendente"] != x["entrada_pendente"] + x["devolvidas"]:
                sys.exit("%s: entrada pendente distribuída não bate com pendente + devolvidas" % x["sku_ml"])
            if x["sku_ml"] not in MAPA:
                sys.exit("SKU do ML sem correspondência no catálogo: %s" % x["sku_ml"])
            linhas.append(x)
    return linhas


def gerar(transcricao, dir_itens, saida, data_txt):
    cat = {}
    for f in glob.glob(os.path.join(dir_itens, "*.json")):
        it = json.load(open(f, encoding="utf-8"))
        it = it.get("data", it)
        cat[it["sku"]] = it

    reg = []
    for l in ler_transcricao(transcricao):
        it = cat[MAPA[l["sku_ml"]]]
        x = dict(l, codigo=it["sku"], custo=it.get("custo") or 0, preco=it.get("preco") or 0, galpao=it.get("qtd") or 0)
        x["full"] = x["aptas"] + x["nao"]
        x["caminho"] = x["entrada_pendente"] + x["em_transferencia"] + x["devolvidas"]
        x["total"] = x["full"] + x["caminho"]
        assert x["total"] == x["ocupam_espaco"]
        reg.append(x)
    reg.sort(key=lambda x: (-x["total"] * x["custo"], x["sku_ml"]))

    galpao = dict(un=sum(it.get("qtd") or 0 for it in cat.values()),
                  custo=round(sum((it.get("qtd") or 0) * (it.get("custo") or 0) for it in cat.values()), 2),
                  rec=round(sum((it.get("qtd") or 0) * (it.get("preco") or 0) for it in cat.values()), 2))

    cache = {"Resumo": {}, "Full por SKU": {}}
    wb = Workbook()

    # ------------------------------ Full por SKU ------------------------------
    det = wb.active
    det.title = "Full por SKU"
    colunas = [("sku", "SKU no ML", 13), ("cod", "Código", 11), ("prod", "Produto", 44),
               ("apt", "Aptas para venda", 10), ("nao", "Não aptas", 8), ("full", "Total no Full", 9),
               ("pen", "Entrada pendente (em contagem)", 12), ("trf", "Em transferência (a caminho)", 12),
               ("dev", "Devolvidas pelo comprador", 11), ("cam", "Total pendente / a caminho", 11),
               ("tot", "Ocupam espaço no Full (ML)", 11), ("cus", "Custo unit. (R$)", 11),
               ("pre", "Preço venda líquido (R$)", 11), ("cpar", "Custo total (R$)", 14),
               ("rec", "Receita projetada (R$)", 14), ("mar", "Margem projetada (R$)", 14),
               ("gal", "No seu galpão (app)", 10)]
    C = {k: i for i, (k, _, _) in enumerate(colunas, 1)}
    L = {k: get_column_letter(v) for k, v in C.items()}
    for k, t, w in colunas:
        cel = det.cell(row=1, column=C[k], value=t)
        cel.font, cel.fill = fonte(True, BRANCO), PatternFill("solid", fgColor=MARCA)
        cel.alignment = Alignment(vertical="center", wrap_text=True)
        det.column_dimensions[L[k]].width = w
    det.row_dimensions[1].height = 44
    det.freeze_panes = "D2"

    L0, ULT = 2, 1 + len(reg)
    TOT = ULT + 1
    inteiros = {C[k] for k in ("apt", "nao", "full", "pen", "trf", "dev", "cam", "tot", "gal")}
    dinheiro = {C[k] for k in ("cus", "pre", "cpar", "rec", "mar")}
    cd = cache["Full por SKU"]

    for i, x in enumerate(reg):
        r = L0 + i
        for k, v in (("sku", x["sku_ml"]), ("cod", x["codigo"]), ("prod", x["produto"]), ("apt", x["aptas"]),
                     ("nao", x["nao"]), ("pen", x["entrada_pendente"]), ("trf", x["em_transferencia"]),
                     ("dev", x["devolvidas"]), ("cus", x["custo"]), ("pre", x["preco"]), ("gal", x["galpao"])):
            det.cell(row=r, column=C[k], value=v)

        def f(k, frm, val):
            det.cell(row=r, column=C[k], value=frm)
            cd["%s%d" % (L[k], r)] = val
        f("full", "=%s%d+%s%d" % (L["apt"], r, L["nao"], r), x["full"])
        f("cam", "=%s%d+%s%d+%s%d" % (L["pen"], r, L["trf"], r, L["dev"], r), x["caminho"])
        f("tot", "=%s%d+%s%d" % (L["full"], r, L["cam"], r), x["total"])
        f("cpar", "=%s%d*%s%d" % (L["tot"], r, L["cus"], r), round(x["total"] * x["custo"], 2))
        f("rec", "=%s%d*%s%d" % (L["tot"], r, L["pre"], r), round(x["total"] * x["preco"], 2))
        f("mar", "=%s%d-%s%d" % (L["rec"], r, L["cpar"], r), round(x["total"] * (x["preco"] - x["custo"]), 2))
        for col in range(1, len(colunas) + 1):
            cel = det.cell(row=r, column=col)
            cel.font, cel.border = fonte(), BORDA
            if col in inteiros: cel.number_format = "#,##0"
            if col in dinheiro: cel.number_format = MOEDA
        for k in ("cus", "pre"):
            det.cell(row=r, column=C[k]).fill = PatternFill("solid", fgColor=AMARELO)
            det.cell(row=r, column=C[k]).font = fonte(cor=AZUL)

    campo = {"apt": "aptas", "nao": "nao", "full": "full", "pen": "entrada_pendente", "trf": "em_transferencia",
             "dev": "devolvidas", "cam": "caminho", "tot": "total", "gal": "galpao"}
    det.cell(row=TOT, column=1, value="TOTAL")
    somas = {k: sum(x[c] for x in reg) for k, c in campo.items()}
    for k in ("cpar", "rec", "mar"):
        somas[k] = round(sum(cd["%s%d" % (L[k], r)] for r in range(L0, ULT + 1)), 2)
    for k, v in somas.items():
        det.cell(row=TOT, column=C[k], value="=SUM({c}{a}:{c}{b})".format(c=L[k], a=L0, b=ULT))
        cd["%s%d" % (L[k], TOT)] = v
    for col in range(1, len(colunas) + 1):
        cel = det.cell(row=TOT, column=col)
        cel.font, cel.fill = fonte(True), PatternFill("solid", fgColor=CINZA)
        cel.border = Border(top=Side(style="medium", color=MARCA))
        if col in inteiros: cel.number_format = "#,##0"
        if col in dinheiro: cel.number_format = MOEDA
    det.auto_filter.ref = "A1:%s%d" % (L["gal"], ULT)
    det.sheet_view.showGridLines = False

    # ------------------------------ Resumo ------------------------------
    res = wb.create_sheet("Resumo", 0)
    cr = cache["Resumo"]
    for col, w in zip("ABCDEF", (40, 16, 18, 20, 18, 60)):
        res.column_dimensions[col].width = w
    res["A1"] = "Mercado Livre Full — posição de %s" % data_txt
    res["A1"].font = fonte(True, MARCA, 15)
    res["A2"] = ("Fonte: print do relatório de estoque Full do Mercado Livre + custo e preço de venda líquido "
                 "do seu catálogo (app Estoque Vivo).")
    res["A2"].font = fonte(cor=SUAVE, italic=True)

    D = "'Full por SKU'!"
    rng = lambda k: "%s%s%d:%s%d" % (D, L[k], L0, L[k], ULT)

    def estilo(r, negrito=False, cinza=False):
        for col in range(1, 7):
            cel = res.cell(row=r, column=col)
            cel.font = fonte(negrito, cor=(SUAVE if col == 6 else None), size=(9 if col == 6 else 10))
            cel.border = BORDA
            if cinza: cel.fill = PatternFill("solid", fgColor=CINZA)
            if col == 6: cel.alignment = Alignment(wrap_text=True, vertical="top")
        res.cell(row=r, column=2).number_format = "#,##0"
        for col in (3, 4, 5):
            res.cell(row=r, column=col).number_format = MOEDA

    def cab(r, titulos):
        for i, t in enumerate(titulos):
            cel = res.cell(row=r, column=1 + i, value=t)
            cel.font, cel.fill = fonte(True, BRANCO), PatternFill("solid", fgColor=MARCA)
            cel.alignment = Alignment(vertical="center", wrap_text=True)
        res.row_dimensions[r].height = 30

    def put(ref, formula, valor):
        res[ref] = formula
        cr[ref] = valor

    def linha(r, rotulo, k, obs):
        c = campo[k]
        un = sum(x[c] for x in reg)
        custo = round(sum(x[c] * x["custo"] for x in reg), 2)
        rec = round(sum(x[c] * x["preco"] for x in reg), 2)
        res.cell(row=r, column=1, value=rotulo)
        put("B%d" % r, "=SUM(%s)" % rng(k), un)
        put("C%d" % r, "=SUMPRODUCT(%s,%s)" % (rng(k), rng("cus")), custo)
        put("D%d" % r, "=SUMPRODUCT(%s,%s)" % (rng(k), rng("pre")), rec)
        put("E%d" % r, "=D%d-C%d" % (r, r), round(rec - custo, 2))
        res.cell(row=r, column=6, value=obs)
        estilo(r)

    def soma(r, rotulo, linhas_, obs):
        res.cell(row=r, column=1, value=rotulo)
        for col in "BCDE":
            put("%s%d" % (col, r), "+".join("%s%d" % (col, x) for x in linhas_).join(("=", "")),
                round(sum(cr["%s%d" % (col, x)] for x in linhas_), 2))
        res.cell(row=r, column=6, value=obs)
        estilo(r, True, True)

    cab(4, ["Situação no Mercado Livre", "Unidades", "Custo (R$)", "Receita projetada (R$)",
            "Margem projetada (R$)", "O que significa"])
    linha(5, "Aptas para venda", "apt", "Já estão no Full e à venda.")
    linha(6, "Temporariamente não aptas", "nao", "Extraviadas, em revisão ou de vendas canceladas.")
    soma(7, "TOTAL NO FULL", (5, 6), "Dentro dos galpões do ML.")
    linha(8, "Entrada pendente (em contagem)", "pen", "Já chegaram ao Full e o ML ainda está conferindo.")
    linha(9, "Em transferência (a caminho)", "trf", "Envios que saíram do seu galpão e ainda não chegaram ao Full.")
    linha(10, "Devolvidas pelo comprador", "dev",
          "Devoluções voltando ao Full. Podem voltar avariadas: o valor é o máximo.")
    soma(11, "TOTAL PENDENTE / A CAMINHO", (8, 9, 10), "Ainda não estão à venda.")
    soma(12, "TOTAL QUE OCUPA ESPAÇO NO FULL", (7, 11), "É o número que o ML mostra na coluna de espaço ocupado.")

    res.cell(row=14, column=1, value="Visão da empresa").font = fonte(True, MARCA, 12)
    cab(15, ["Onde está", "Unidades", "Custo (R$)", "Receita projetada (R$)", "Margem projetada (R$)", "Observação"])
    res["A16"] = "Seu galpão (app Estoque Vivo)"
    res["B16"], res["C16"], res["D16"] = galpao["un"], galpao["custo"], galpao["rec"]
    put("E16", "=D16-C16", round(galpao["rec"] - galpao["custo"], 2))
    res["F16"] = "Números do app em %s. Os envios em transferência já saíram do galpão no app." % data_txt
    estilo(16)
    for col in "BCD":
        res["%s16" % col].font = fonte(cor=AZUL)
        res["%s16" % col].comment = Comment("Número copiado do app Estoque Vivo (não é fórmula).", "Estoque Vivo")
    res["A17"] = "Mercado Livre Full (no Full + a caminho)"
    for col in "BCDE":
        put("%s17" % col, "=%s12" % col, cr["%s12" % col])
    res["F17"] = "Total que ocupa espaço no Full, da tabela acima."
    estilo(17)
    res["A18"] = "TOTAL DA EMPRESA"
    for col, g in zip("BCDE", (galpao["un"], galpao["custo"], galpao["rec"], galpao["rec"] - galpao["custo"])):
        put("%s18" % col, "=%s16+%s17" % (col, col), round(g + cr["%s17" % col], 2))
    res["F18"] = "Tudo o que é seu: galpão, Full, a caminho e devoluções."
    estilo(18, True, True)

    notas = [
        "Receita projetada = unidades × seu preço de venda líquido (o que você recebe por peça, já sem as tarifas do ML). "
        "Margem = receita − custo.",
        "Custo e preço vêm do seu catálogo (células amarelas da aba Full por SKU). Se mudar um valor ali, o resumo recalcula sozinho.",
        "Não aptas = espaço ocupado − (aptas + entrada pendente + em transferência + devolvidas). O print esconde as "
        "colunas de não aptas do ML; a conta fecha com o total de %s unidades do relatório." % format(somas["tot"], ",").replace(",", "."),
        "O print não traz vendas dos últimos 30 dias nem urgência de envio. Para ter cobertura em dias, mande o arquivo "
        ".xlsx do relatório do ML.",
    ]
    res.cell(row=20, column=1, value="Como ler").font = fonte(True, MARCA, 12)
    for r, t in enumerate(notas, 21):
        cel = res.cell(row=r, column=1, value="• " + t)
        cel.font, cel.alignment = fonte(), Alignment(wrap_text=True, vertical="top")
        res.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        res.row_dimensions[r].height = 30
    res.sheet_view.showGridLines = False

    wb.save(saida)
    injetar_cache(saida, cache)
    validar(saida, cache)
    return cache


if __name__ == "__main__":
    if len(sys.argv) != 5:
        sys.exit(__doc__)
    r = gerar(*sys.argv[1:])["Resumo"]
    for rot, lin in (("No Full", 7), ("Pendente/a caminho", 11), ("Ocupa espaço", 12), ("Empresa", 18)):
        print("%-19s %5d un | custo R$ %10.2f | receita R$ %10.2f" % (rot, r["B%d" % lin], r["C%d" % lin], r["D%d" % lin]))
